"""集計と、書式付き報告帳票（xlsx）の出力。"""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_crosstab(df: pd.DataFrame) -> pd.DataFrame:
    """支店 × 種別のクロス集計を作る。合計行・列を付ける。

    元ファイルの合計行は読み込み時に除去済みで、ここで明細から計算し直している。
    数式セルの値を当てにしないため（Excel で保存されていないファイルでは None になる）。
    """
    table = pd.pivot_table(
        df, index="支店", columns="種別", values="件数",
        aggfunc="sum", fill_value=0, observed=True,
    )
    table["合計"] = table.sum(axis=1)
    table.loc["合計"] = table.sum(axis=0)
    return table


def write_report(table: pd.DataFrame, path: Path, target_month: date, note_lines: list[str]) -> None:
    """書式付きの帳票を書き出す。

    数式ではなく計算済みの値を書く。受け取った側が Excel で開かなくても
    値が読める状態にするため（この帳票自体が別のバッチの入力になりうる）。
    """
    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="集計", startrow=2)
        ws = writer.sheets["集計"]

        ws["A1"] = f"支店別 問い合わせ集計　{target_month.year}年{target_month.month}月"
        ws["A1"].font = Font(bold=True, size=14)

        n_rows, n_cols = table.shape
        header_row = 3
        first_data_row = header_row + 1
        last_data_row = header_row + n_rows

        for col in range(1, n_cols + 2):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for row in range(first_data_row, last_data_row + 1):
            for col in range(1, n_cols + 2):
                cell = ws.cell(row=row, column=col)
                cell.border = BORDER
                if col > 1:
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right")
            # 合計行と合計列に色を付ける
            if row == last_data_row:
                for col in range(1, n_cols + 2):
                    ws.cell(row=row, column=col).fill = TOTAL_FILL
                    ws.cell(row=row, column=col).font = Font(bold=True)
        for row in range(first_data_row, last_data_row + 1):
            ws.cell(row=row, column=n_cols + 1).fill = TOTAL_FILL
            ws.cell(row=row, column=n_cols + 1).font = Font(bold=True)

        ws.column_dimensions["A"].width = 12
        for col in range(2, n_cols + 2):
            ws.column_dimensions[get_column_letter(col)].width = 13

        # 読み込みの内訳を同じブックに残す。
        # 「何行除去したか」を帳票の受け取り側が確認できるようにするため。
        ws_note = writer.book.create_sheet("読み込み内訳")
        ws_note["A1"] = "読み込み内訳（この帳票がどのデータから作られたか）"
        ws_note["A1"].font = Font(bold=True)
        for i, line in enumerate(note_lines, start=3):
            ws_note.cell(row=i, column=1, value=line)
        ws_note.column_dimensions["A"].width = 100
