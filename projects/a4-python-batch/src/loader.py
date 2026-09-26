"""支店別ファイルの読み込みと正規化。

【設計方針】黙って直さない。
変換できなかった値・除去した行は件数を記録し、呼び出し側が確認できるようにする。
バッチ処理は「エラーにならず中身が違う」形で壊れるため、何を捨てたかが見えないと検知できない。

観察した壊れ方と対処は README を参照。
"""

import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

# ヘッダ行の判定に使う語。この3つが同じ行に揃っていればヘッダとみなす。
HEADER_MARKERS = ("日付", "支店", "種別")

# 明細行から除外する種別の値（元ファイルの合計行）
SUMMARY_LABELS = {"合計", "総計", "小計"}

# 読み込みを試す文字コード。上から順に試し、最初に成功したものを使う。
# 推測で読み進めず、全部失敗したらエラーにする。
ENCODINGS = ("utf-8-sig", "cp932", "utf-8")


@dataclass
class LoadReport:
    """読み込み結果の内訳。何を捨てたかを呼び出し側に見せるために持つ。"""

    source: str
    encoding: str | None = None
    header_row: int | None = None
    total_rows: int = 0
    detail_rows: int = 0
    dropped_empty: int = 0
    dropped_summary: int = 0
    invalid_dates: list[str] = field(default_factory=list)
    invalid_counts: list[str] = field(default_factory=list)

    def summary(self) -> str:
        """1行の要約。ログに出す用。"""
        parts = [f"{self.source}: {self.detail_rows}行"]
        if self.encoding:
            parts.append(f"enc={self.encoding}")
        if self.header_row is not None:
            parts.append(f"header={self.header_row}行目")
        if self.dropped_empty:
            parts.append(f"空行除去={self.dropped_empty}")
        if self.dropped_summary:
            parts.append(f"合計行除去={self.dropped_summary}")
        if self.invalid_dates:
            parts.append(f"日付変換不可={len(self.invalid_dates)}")
        if self.invalid_counts:
            parts.append(f"件数変換不可={len(self.invalid_counts)}")
        return " / ".join(parts)


def normalize_text(value: object) -> str:
    """全角・半角・空白のゆれを吸収する。

    NFKC 正規化で「ｼｽﾃﾑ障害」→「システム障害」、「１４」→「14」に揃う。
    全角スペース（U+3000）は NFKC では半角スペースになるだけなので、別途除去する
    （「大 阪」→「大阪」）。
    """
    text = unicodedata.normalize("NFKC", str(value))
    return text.replace(" ", "").strip()


def parse_date(value: object) -> date | None:
    """日付を date に統一する。変換できなければ None を返す（呼び出し側が記録する）。"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_text(value)
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    # pandas に最後の判断を委ねる（"2026-08-02 00:00:00" のような文字列を拾う）
    parsed = pd.to_datetime(text, errors="coerce")
    return None if pd.isna(parsed) else parsed.date()


def parse_count(value: object) -> int | None:
    """件数を int に統一する。全角数字も受ける。変換できなければ None。"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int,)) and not isinstance(value, bool):
        return int(value)
    if isinstance(value, float):
        return int(value) if float(value).is_integer() else None
    text = normalize_text(value)
    return int(text) if text.isdigit() else None


def detect_encoding(path: Path) -> str:
    """CSV の文字コードを判定する。

    候補を順に試し、最初にデコードできたものを採用する。
    **全部失敗したら例外を投げる。** 推測で読み進めると文字化けしたまま処理が進む。
    """
    raw = path.read_bytes()
    for enc in ENCODINGS:
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    raise ValueError(
        f"{path.name} の文字コードを判定できません。試した候補: {', '.join(ENCODINGS)}"
    )


def find_header_row(rows: list[list]) -> int:
    """ヘッダ行の位置（0始まり）を探す。

    結合セルのタイトルが1行目にあるため、固定位置を前提にできない。
    HEADER_MARKERS が揃う行をヘッダとみなす。**見つからなければ例外。**
    """
    for i, row in enumerate(rows[:20]):
        values = {normalize_text(v) for v in row if v is not None}
        if all(marker in values for marker in HEADER_MARKERS):
            return i
    raise ValueError(
        f"ヘッダ行が見つかりません（先頭20行に {', '.join(HEADER_MARKERS)} が揃う行がない）"
    )


def _rows_from_xlsx(path: Path) -> list[list]:
    """xlsx から生の値を読む。

    data_only は指定しない。数式セルは文字列として返るが、
    **合計行は明細から自分で計算するので数式の値は使わない。**
    data_only=True にすると、Excel で保存されていないファイルでは None になる。
    """
    ws = load_workbook(path, data_only=False).active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _rows_from_csv(path: Path, encoding: str) -> list[list]:
    """CSV から生の値を読む。"""
    text = path.read_bytes().decode(encoding)
    return [line.split(",") for line in text.splitlines()]


def load_branch_file(path: Path) -> tuple[pd.DataFrame, LoadReport]:
    """1ファイルを読み、明細行だけの DataFrame と内訳を返す。"""
    report = LoadReport(source=path.name)

    if path.suffix == ".csv":
        report.encoding = detect_encoding(path)
        rows = _rows_from_csv(path, report.encoding)
    elif path.suffix == ".xlsx":
        rows = _rows_from_xlsx(path)
    else:
        raise ValueError(f"対応していない形式です: {path.name}")

    header_index = find_header_row(rows)
    report.header_row = header_index + 1
    headers = [normalize_text(v) for v in rows[header_index]]
    body = rows[header_index + 1 :]
    report.total_rows = len(body)

    records = []
    for row in body:
        cells = dict(zip(headers, row))

        # 空行を除去
        if all(v is None or str(v).strip() == "" for v in row):
            report.dropped_empty += 1
            continue

        category = normalize_text(cells.get("種別", ""))

        # 元ファイルの合計行を除去（数式の値は使わず、明細から計算し直す）
        if category in SUMMARY_LABELS:
            report.dropped_summary += 1
            continue

        parsed_date = parse_date(cells.get("日付"))
        parsed_count = parse_count(cells.get("件数"))

        # 変換できなかった値は記録する。黙って捨てない。
        if parsed_date is None:
            report.invalid_dates.append(str(cells.get("日付")))
        if parsed_count is None:
            report.invalid_counts.append(str(cells.get("件数")))

        records.append(
            {
                "日付": parsed_date,
                "支店": normalize_text(cells.get("支店", "")),
                "種別": category,
                "件数": parsed_count,
                "担当者": normalize_text(cells.get("担当者", "")),
                "元ファイル": path.name,
            }
        )

    report.detail_rows = len(records)
    return pd.DataFrame(records), report


def load_all(input_dir: Path) -> tuple[pd.DataFrame, list[LoadReport]]:
    """ディレクトリ内の全ファイルを読んで結合する。"""
    frames, reports = [], []
    for path in sorted(input_dir.iterdir()):
        if path.suffix not in (".xlsx", ".csv"):
            continue
        df, report = load_branch_file(path)
        frames.append(df)
        reports.append(report)

    if not frames:
        raise ValueError(f"読み込めるファイルがありません: {input_dir}")

    return pd.concat(frames, ignore_index=True), reports
